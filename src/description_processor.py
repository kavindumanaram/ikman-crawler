
from google.oauth2.service_account import Credentials
from bs4 import BeautifulSoup
import time
from openpyxl import load_workbook

from utils import get_spreadsheet, extract_json_data
# from utils import get_spreadsheet, extract_json_data



def get_last_house_data(sheet_name, file_path):
    # old_values = sheet.get_all_values()
    # old_data_obj = []
    # headers = old_values[0]
    # url_index = headers.index('URL')
    # for i in range(1, len(old_values)):
    #     data_row = old_values[i]
    #     url = data_row[url_index]
    #     if url not in old_data_obj:
    #         old_data_obj.append(url)
    # print(f"Last house data processed successfully. Count: {len(old_data_obj)}")
    # return old_data_obj
        # Load the workbook and select the sheet
    try:
        wb = load_workbook(file_path)
        # if sheet_name not in wb.sheetnames:
        #     raise ValueError(f"Sheet '{sheet_name}' does not exist in the workbook.")
        ws = wb["Sale"]
    except FileNotFoundError:
        raise FileNotFoundError(f"The file {file_path} does not exist.")

    old_data_obj = []
    headers = [cell.value for cell in ws[1]]  # Read the header row
    url_index = headers.index('URL')  # Get the index of the 'URL' column

    # Iterate over rows starting from the second row
    for row in ws.iter_rows(min_row=2, values_only=True):
        url = row[url_index]  # Get the URL value from the corresponding column
        if url and url not in old_data_obj:  # Ensure URL is not None and is unique
            old_data_obj.append(url)

    print(f"Last house data processed successfully. Count: {len(old_data_obj)}")
    return old_data_obj
 

def get_last_desc_urls(sheet_name,file_path):
    # old_values = sheet.get_all_values()
    # headers = old_values[0]
    # url_index = headers.index('URL')
    # old_data_urls = []
    # for i in range(1, len(old_values)):
    #     data_row = old_values[i]
    #     url = data_row[url_index]
    #     old_data_urls.append(url)
    # print("Last description data processed successfully")
    # return old_data_urls
        # Load the workbook and select the sheet
    try:
        wb = load_workbook(file_path)
        # if sheet_name not in wb.sheetnames:
        #     raise ValueError(f"Sheet '{sheet_name}' does not exist in the workbook.")
        ws = wb["Description"]
    except FileNotFoundError:
        raise FileNotFoundError(f"The file {file_path} does not exist.")

    old_data_urls = []
    headers = [cell.value for cell in ws[1]]  # Read the header row
    url_index = headers.index('URL')  # Get the index of the 'URL' column

    # Iterate over rows starting from the second row
    for row in ws.iter_rows(min_row=2, values_only=True):
        url = row[url_index]  # Get the URL value from the corresponding column
        if url:  # Ensure URL is not None
            old_data_urls.append(url)

    print("Last description data processed successfully")
    return old_data_urls
 

def append_desc_data(file_path, sheet, data):
    print(f"data={data}")
    # last_row = len(sheet.get_all_values()) + 1
    # required_rows = last_row + len(data)
    # if required_rows > sheet.row_count:
    #     sheet.add_rows(required_rows - sheet.row_count)
    # range_str = 'A' + str(last_row) + ':C' + str(required_rows)
    # sheet.update(range_str, data)
    # print(f"Appended data to : {sheet.title}, count: {len(data)}")
    try:
        wb = load_workbook(file_path)
        # Check if the sheet exists
        #print(f"wb.sheetnames={wb.sheetnames}")
        #if sheet_name in wb.sheetnames:
        ws = wb["Description"]
        #else:
            # Create the sheet if it doesn't exist
        #ws = wb.create_sheet(sheet_name)
    except FileNotFoundError:
        raise FileNotFoundError(f"The file {file_path} does not exist.")

    # Append each row of data to the sheet
    for row_data in data:
        ws.append(row_data)

    # Save the workbook
    wb.save(file_path)
    print(f"Appended data to sheet '{sheet}' in file '{file_path}'.")

def handler(event, context):
    spreadsheet = get_spreadsheet()
    house_data_sheet = spreadsheet.parse(sheet_name="Sale")  # Load "Sale" sheet into a DataFrame
    last_house_data = get_last_house_data(house_data_sheet, "test.xlsx")
    print(f"last_house_data={last_house_data}")
    
    desc_data_sheet = spreadsheet.parse(sheet_name="Description") #spreadsheet.worksheet("Description")
    desc_data_urls = get_last_desc_urls(desc_data_sheet, "test.xlsx")     
    print(f"desc_data_urls={desc_data_urls}")
    return_data = []
        
    start_time = time.time()
    for url in last_house_data:        
        if time.time() - start_time > 810:  # 13.5 minutes in seconds
            break
        try:
            if url not in desc_data_urls:
                data = extract_json_data(url)
                #print(f"data url={data}")
                desc = data['adDetail']['data']['ad']['description']

                properties = data['adDetail']['data']['ad']['properties']
                land_size = next((property['value'].split()[0] for property in properties if property['key'] == "land_size"), None)
                return_data.append([
                    f"{url}",
                    desc,
                    land_size
                ])                    
        except Exception as e:
            print(f"Error occurred: {e}. Skipping this iteration.")
            continue
        
     

    append_desc_data("test.xlsx",desc_data_sheet, return_data)
    
    
handler({}, {})


